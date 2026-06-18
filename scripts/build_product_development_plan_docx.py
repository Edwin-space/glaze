from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "04_product_development_plan.docx"


FONT = "Apple SD Gothic Neo"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(80, 80, 80)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9E2EC"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def add_title(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=4, line=1.0)
    run = p.add_run("글레이즈 기본 제품/개발 계획")
    set_run_font(run, size=24, bold=True, color=RGBColor(11, 37, 69))

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=14)
    r = p.add_run("벤치마크, 제품 방향, 기능 범위, 디자인 원칙, 개발 단계 정리")
    set_run_font(r, size=11, color=MUTED)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=12)
    r = p.add_run("작성일: 2026-06-18  |  프로젝트명: 글레이즈")
    set_run_font(r, size=10, color=MUTED)


def add_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, after=4)
        r = p.add_run(item)
        set_run_font(r)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_spacing(p, after=4)
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_table_borders(table, color="D8DEE8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=2)
    r = p.add_run(label)
    set_run_font(r, size=10, bold=True, color=DARK_BLUE)
    p = cell.add_paragraph()
    set_paragraph_spacing(p, after=0)
    r = p.add_run(text)
    set_run_font(r, size=11)


def add_competitor_table(doc):
    rows = [
        ("Infuse", "Apple 생태계형 미디어 라이브러리 플레이어의 기준점", "라이브러리, NAS/서버 연동, 메타데이터 UX", "초기부터 모든 서버/코덱을 따라잡으려 하지 않는다"),
        ("VidHub", "Infuse 대안 포지션의 라이브러리 플레이어", "개인 영상 보관함 중심 구조와 자막 상태 관리", "AI 자막 준비 경험은 별도 차별화로 설계한다"),
        ("Submarine Player", "AI 자막 생성과 번역을 이미 제품화한 직접 레퍼런스", "실시간 자막, 오프라인 처리, 이중 자막, 배치 처리, 내보내기", "글레이즈는 한국어 중심 사전 생성과 라이브러리 자동화로 차별화한다"),
        ("MX Video Player HD", "리뷰 수가 큰 범용 플레이어", "대중적 포맷 지원 요구 확인", "단순 포맷 지원 앱처럼 보이지 않게 한다"),
        ("Movist 등", "한국 사용자 관점에서 참고할 고급 플레이어", "자막 품질, 디코더 선택, 세부 설정", "MVP 범위를 흐리지 않게 관찰 대상으로 둔다"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    headers = ["제품", "경쟁 의미", "배울 점", "초기 판단"]
    widths = [1500, 2600, 2600, 2660]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Pt(widths[i] / 20)
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0)
        r = p.add_run(h)
        set_run_font(r, size=10, bold=True, color=RGBColor(40, 40, 40))
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Pt(widths[i] / 20)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.08)
            r = p.add_run(val)
            set_run_font(r, size=9.5, bold=(i == 0))


def add_phase_table(doc):
    rows = [
        ("0", "프로젝트 기반 정리", "Git 저장소, Swift/macOS 프로젝트, 빌드/실행 기준 정리"),
        ("1", "플레이어 MVP", "로컬 파일 열기, 재생 컨트롤, 전체화면, 외부 자막 로딩"),
        ("2", "수동 자막 생성 MVP", "오디오 추출, 음성 인식, SRT 저장, 자동 연결, 진행률"),
        ("3", "한국어 번역 자막", "문장 단위 번역, 한국어 SRT 저장, 이중 자막 표시"),
        ("4", "자막 준비 큐", "재생목록 감지, 작업 상태 저장, 재시도, 일시정지/재개"),
        ("5", "라이브러리와 자동화", "폴더 등록, 상태 표시, 전원/배터리 정책, 백그라운드 처리"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    headers = ["Phase", "목표", "주요 산출물"]
    widths = [900, 2400, 6060]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Pt(widths[i] / 20)
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0)
        r = p.add_run(h)
        set_run_font(r, size=10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Pt(widths[i] / 20)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, after=0, line=1.08)
            r = p.add_run(val)
            set_run_font(r, size=9.5, bold=(i == 0))


def add_platform_table(doc):
    rows = [
        ("macOS Apple Silicon", "초기 주력", "AI 자막 생성, 번역, 라이브러리, 배치 처리"),
        ("iOS/iPadOS", "보조 감상/동기화", "준비된 자막 영상 재생, 학습, 짧은 영상 생성"),
        ("Android", "글로벌 모바일 확장 후보", "모바일 감상, 학습, 경량 생성"),
        ("Windows", "글로벌 데스크톱 확장 후보", "데스크톱 재생, 배치 생성, 라이브러리"),
        ("브라우저 확장", "웹 영상 브릿지", "DRM 없는 웹 영상 열기, HLS 감지"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    headers = ["플랫폼", "역할", "핵심 기능"]
    widths = [2200, 2600, 4560]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Pt(widths[i] / 20)
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0)
        r = p.add_run(h)
        set_run_font(r, size=10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Pt(widths[i] / 20)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.08)
            r = p.add_run(val)
            set_run_font(r, size=9.5, bold=(i == 0))


def build():
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("1. 문서 목적", level=1)
    add_p(doc, "이 문서는 글레이즈를 실제 개발 단계로 진입시키기 위한 기본 계획서다. 앞서 정리한 Mac App Store 벤치마크, 제품 방향, AI 자막 생성 워크플로우를 하나의 실행 기준으로 묶고, 기능 범위, 디자인 방향, 기술 방향, MVP 개발 순서를 정의한다.")

    doc.add_heading("2. 제품 판단", level=1)
    add_p(doc, "글레이즈는 범용 영상 플레이어 시장에서 Infuse, VidHub, MX Video Player HD 같은 제품과 직접 경쟁한다. 다만 초기 전략은 가장 많은 코덱을 지원하는 플레이어가 아니라 자막 없는 영상을 감상 가능한 상태로 준비해주는 AI 플레이어가 되어야 한다.")
    add_callout(doc, "핵심 포지션", "Apple Silicon 온디바이스 AI로 자막 없는 다국어 영상을 재생 전에 자막화하고, 한국어 감상 경험까지 준비해주는 macOS 영상 플레이어.")
    add_p(doc, "2차 차별화 후보로는 Chrome/Safari 브라우저 확장 연동을 둔다. 사용자가 웹에서 보던 DRM 없는 영상을 글레이즈로 열고, 자체 플레이어에서 자막 생성과 번역을 이어갈 수 있게 하는 방향이다.")

    doc.add_heading("3. 벤치마크 기반 경쟁 구도", level=1)
    add_competitor_table(doc)

    doc.add_heading("4. 차별화 원칙", level=1)
    add_bullets(doc, [
        "기존 플레이어는 있는 자막을 잘 붙이는 쪽에 강하다.",
        "글레이즈는 없는 자막을 만들어 재사용 가능한 자산으로 남기는 쪽에 집중한다.",
        "AI 기능은 전면 장식이 아니라 감상 준비를 자동화하는 기반 기능이어야 한다.",
        "한국어 번역 자막은 부가 기능이 아니라 기본 경험에 가깝게 설계한다.",
        "재생 중 실시간 처리보다 재생 전 사전 생성과 백그라운드 큐를 우선한다.",
        "Submarine Player가 이미 보여준 실시간 AI 자막 경험을 전제로, 글레이즈는 라이브러리 준비형 AI 플레이어로 포지셔닝한다.",
        "브라우저 확장 연동은 보호 콘텐츠 우회 도구가 아니라, 정당하게 접근 가능한 웹 영상을 글레이즈 감상 환경으로 가져오는 브릿지로 설계한다.",
    ])

    doc.add_heading("5. MVP 기능 범위", level=1)
    add_p(doc, "MVP는 영상 재생, 단일 영상 수동 AI 자막 생성, 저장, 자동 연결, 한국어 번역까지를 목표로 한다.")
    doc.add_heading("포함 기능", level=2)
    add_bullets(doc, [
        "macOS 네이티브 앱과 로컬 영상 파일 열기",
        "기본 재생, 일시정지, 탐색, 전체화면",
        "외부 자막 파일 로딩",
        "단일 영상의 오디오 추출과 온디바이스 음성 인식",
        "원어 자막 생성, 저장, 자동 연결",
        "한국어 번역 자막 생성과 원문/한국어/이중 자막 보기",
        "기본 자막 스타일 설정과 자막 생성 진행 상태 표시",
    ])
    doc.add_heading("초기 제외 기능", level=2)
    add_bullets(doc, [
        "초기부터 모든 NAS, 클라우드, 미디어 서버 연동",
        "실시간 자막 생성 중심 UX",
        "복잡한 라이브러리 자동 메타데이터 매칭",
        "모든 코덱 완전 지원",
        "AI 요약, 챕터, 학습 모드의 완성형 제공",
    ])

    doc.add_heading("6. 디자인 진행 방향", level=1)
    add_p(doc, "macOS 네이티브 감각을 우선한다. 영상 감상 화면은 조용하고 어둡게 유지하고, 설정과 라이브러리 화면은 밝고 정보 중심으로 설계한다. AI 용어보다 사용자가 얻는 결과를 먼저 보여주며, 상태와 실패 사유는 숨기지 않는다.")
    doc.add_heading("핵심 화면", level=2)
    add_bullets(doc, [
        "플레이어 화면: 영상 영역, 재생 컨트롤, 탐색바, 자막 토글, 자막 언어 선택, 전체화면과 PIP",
        "자막 생성 패널: 상태, 진행률, 예상 시간, 모델 모드, 출력 옵션, 저장 위치, 작업 제어",
        "라이브러리 화면: 영상 목록, 자막 상태, 번역 상태, 마지막 재생 위치, 작업 큐 상태, 실패 사유",
        "자막 검수 화면: 생성 자막 수정, 싱크 확인, 원문/번역 비교, SRT/VTT/CSV 내보내기",
        "브라우저 확장 화면: 감지된 HLS/미디어 URL, 글레이즈에서 열기, DRM 감지 시 비활성화, 허용된 미디어 저장",
    ])

    doc.add_heading("7. 기술 방향", level=1)
    add_bullets(doc, [
        "SwiftUI: 앱 UI와 설정 화면",
        "AVFoundation/AVKit: 기본 재생",
        "FFmpeg 검토: MKV, 특수 코덱, 오디오 추출 보완",
        "Core Data 또는 SQLite: 라이브러리, 자막 인덱스, 작업 큐 저장",
        "Core ML/Metal: 온디바이스 음성 인식 모델 실행 검토",
        "SRT/VTT: 초기 자막 저장 포맷",
        "브라우저 확장 후보: Chrome Manifest V3, Safari Web Extension, 커스텀 URL 스킴 또는 Native Messaging",
    ])

    doc.add_heading("8. 개발 단계", level=1)
    add_phase_table(doc)

    doc.add_heading("9. 플랫폼 및 글로벌 확장", level=1)
    add_p(doc, "글레이즈는 초기부터 글로벌 제품을 지향하지만 모든 플랫폼을 동시에 지원하지 않는다. 먼저 macOS Apple Silicon에서 핵심 가치를 증명하고, 이후 iOS/iPadOS 보조 앱, Android 검토, Windows 데스크톱 제품군으로 순차 확장한다.")
    add_platform_table(doc)
    doc.add_heading("언어 전략", level=2)
    add_p(doc, "초기 UI 기본 언어는 한국어와 영어로 둔다. 이후 일본어, 중국어, 스페인어, 프랑스어, 독일어 등 주요 국가 언어는 실제 유입과 결제 전환을 기준으로 확장한다. 기계 번역 초안을 활용하되 결제, 개인정보, 다운로드, DRM 관련 문구는 반드시 사람이 검수한다.")

    doc.add_heading("10. 초기 의사결정 필요 사항", level=1)
    add_bullets(doc, [
        "MVP에서 지원할 최소 macOS 버전",
        "초기 영상 재생을 AVFoundation 중심으로 할지, FFmpeg를 조기에 붙일지",
        "음성 인식 모델 후보와 라이선스",
        "한국어 번역을 온디바이스 우선으로 할지, 초기에는 온라인 API 선택 옵션을 둘지",
        "생성 자막의 기본 저장 위치",
        "생성 자막 편집 기능을 MVP에 포함할지, 1차 이후로 둘지",
        "SRT/VTT/CSV 내보내기를 MVP에 포함할지",
        "브라우저 확장 연동을 어느 Phase에서 실험할지",
        "HLS 감지와 재생까지만 할지, 허용된 미디어 저장까지 제공할지",
        "DRM/EME 감지와 차단 정책을 어떻게 사용자에게 설명할지",
        "초기 UI 언어를 한국어/영어로 확정할지",
        "iOS/iPadOS를 보조 감상 앱으로 둘지, 독립 생성 앱으로 키울지",
        "Windows 진입 시 AI 엔진을 어떤 런타임으로 이식할지",
        "앱 이름, 번들 ID, App Store 배포 여부와 샌드박스 정책",
    ])

    doc.add_heading("11. 개발 착수 전 체크리스트", level=1)
    add_bullets(doc, [
        "제품 포지션 문장 확정",
        "MVP 기능 범위와 제외 기능 확정",
        "기술 스택 1차 확정",
        "SwiftUI macOS 프로젝트 생성",
        "Git 저장소 초기화",
        "한국어/영어 localization 구조 적용",
        "기본 화면 와이어프레임 작성",
        "자막 생성 파이프라인 기술 검증",
        "샘플 영상 세트 준비",
        "벤치마크 제품 기능 체크리스트 확장",
    ])

    doc.add_heading("12. 다음 행동", level=1)
    add_p(doc, "다음 단계는 개발 프로젝트의 기반을 만드는 것이다. 우선 SwiftUI 기반 macOS 앱 프로젝트를 생성하고, 기본 영상 파일 열기와 재생 화면을 구현한다. 동시에 AI 자막 생성은 별도 기술 검증 트랙으로 분리해, 오디오 추출과 음성 인식 모델 실행 가능성을 먼저 확인한다.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
